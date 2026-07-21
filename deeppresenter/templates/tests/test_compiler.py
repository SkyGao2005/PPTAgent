from __future__ import annotations

import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pydantic import ValidationError

from deeppresenter.templates.annotation import (
    DeterministicAnnotator,
    SlideAnnotationInput,
)
from deeppresenter.templates.compiler import TemplateCompilationError, TemplateCompiler
from deeppresenter.templates.ids import derive_revision_id
from deeppresenter.templates.models import (
    AssetIndex,
    AssetRole,
    NormalizedBox,
    ReusePolicy,
)
from deeppresenter.templates.rendering import NullRenderer, RenderedSlide
from deeppresenter.templates.store import TemplateStore


def _make_template(path: Path, *, slides: int = 2) -> None:
    logo_path = path.parent / "logo.png"
    background_path = path.parent / "background.png"
    Image.new("RGB", (160, 80), "#e43d30").save(logo_path)
    Image.new("RGB", (1600, 900), "#102040").save(background_path)

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for page_number in range(1, slides + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.shapes.add_picture(
            str(background_path),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        background.name = "Background"
        title = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.6), Inches(10), Inches(0.8)
        )
        title.name = "Title"
        run = title.text_frame.paragraphs[0].add_run()
        run.text = f"Template page {page_number}"
        run.font.name = "Aptos"
        run.font.size = Pt(30)
        logo = slide.shapes.add_picture(
            str(logo_path), Inches(11.6), Inches(0.2), Inches(1), Inches(0.5)
        )
        logo.name = "Company Logo"
    presentation.save(path)


def _make_nonstandard_template(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(12)
    presentation.slide_height = Inches(8)
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)


@pytest.mark.asyncio
async def test_compiler_publishes_content_addressed_assets_and_is_idempotent(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.pptx"
    _make_template(source)
    template_dir = tmp_path / "templates" / "brand"
    compiler = TemplateCompiler(renderer=NullRenderer())

    first = await compiler.compile("brand", source, template_dir, name="Brand")
    second = await compiler.compile("brand", source, template_dir, name="Brand")

    assert first.latest_revision_id == second.latest_revision_id
    assert first.active_revision_id == first.latest_revision_id
    assert len(list((template_dir / "revisions").iterdir())) == 1
    revision_dir = template_dir / "revisions" / first.latest_revision_id
    asset_index = AssetIndex.model_validate_json(
        (revision_dir / "assets" / "index.json").read_text(encoding="utf-8")
    )
    by_role = {asset.role: asset for asset in asset_index.assets}
    assert by_role[AssetRole.LOGO].reuse_policy is ReusePolicy.ALWAYS
    assert by_role[AssetRole.BACKGROUND].reuse_policy is ReusePolicy.TEMPLATE_ONLY
    assert (revision_dir / by_role[AssetRole.LOGO].path).read_bytes()
    assert by_role[AssetRole.LOGO].browser_path == by_role[AssetRole.LOGO].path
    assert by_role[AssetRole.LOGO].browser_media_type == "image/png"
    assert by_role[AssetRole.LOGO].browser_sha256 == by_role[AssetRole.LOGO].sha256

    runtime_revision = TemplateStore(tmp_path / "templates").resolve("brand")
    assert runtime_revision.revision_id == first.latest_revision_id
    assert len(runtime_revision.load_slide_index()) == 2
    assert runtime_revision.load_slide_compact("s001")["slide_id"] == "s001"


@pytest.mark.asyncio
async def test_compiler_rejects_nonstandard_slide_ratio(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    _make_nonstandard_template(source)
    template_dir = tmp_path / "templates" / "nonstandard"

    with pytest.raises(TemplateCompilationError, match="16:9 or 4:3"):
        await TemplateCompiler(renderer=NullRenderer()).compile(
            "nonstandard",
            source,
            template_dir,
        )

    manifest = json.loads((template_dir / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "failed"
    assert not (template_dir / "revisions").exists()


class _FakeRenderer:
    renderer_id = "fake-renderer-v1"

    async def render(
        self,
        source_pptx: Path,
        output_dir: Path,
        slide_count: int,
    ) -> list[RenderedSlide]:
        del source_pptx
        output_dir.mkdir(parents=True, exist_ok=True)
        result = []
        for page_number in range(1, slide_count + 1):
            path = output_dir / f"{page_number}.png"
            Image.new("RGB", (800, 450), "white").save(path)
            result.append(RenderedSlide(page_number=page_number, image_path=path))
        return result


class _RecordingAnnotator:
    annotator_id = "recording-annotator-v1"

    def __init__(self) -> None:
        self.requests: list[SlideAnnotationInput] = []
        self.fallback = DeterministicAnnotator()

    async def annotate(self, request: SlideAnnotationInput):
        self.requests.append(request)
        assert request.source_graph.page_number == len(self.requests)
        assert request.reference_image_path is not None
        assert request.overlay_image_path is not None
        value = await self.fallback.annotate(request)
        return value.model_copy(update={"annotator_id": self.annotator_id})


@pytest.mark.asyncio
async def test_renderer_and_annotator_are_invoked_once_per_slide(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.pptx"
    _make_template(source, slides=2)
    annotator = _RecordingAnnotator()
    compiler = TemplateCompiler(renderer=_FakeRenderer(), annotator=annotator)

    manifest = await compiler.compile("visual", source, tmp_path / "visual")

    assert [request.source_graph.slide_id for request in annotator.requests] == [
        "s001",
        "s002",
    ]
    revision = tmp_path / "visual" / "revisions" / manifest.latest_revision_id
    assert (revision / "slides" / "s001" / "reference.webp").exists()
    assert (revision / "slides" / "s001" / "overlay.webp").exists()
    report = json.loads(
        (revision / "validation" / "report.json").read_text(encoding="utf-8")
    )
    assert report["valid"] is True
    assert report["stats"]["rendered_slides"] == 2


class _FlakyAnnotator:
    """Fails one slide until told otherwise; records every annotate call."""

    annotator_id = "flaky-annotator-v1"

    def __init__(self) -> None:
        self.calls: list[str] = []
        self.fail_slide: str | None = "s002"
        self.fallback = DeterministicAnnotator()

    async def annotate(self, request: SlideAnnotationInput):
        self.calls.append(request.source_graph.slide_id)
        if request.source_graph.slide_id == self.fail_slide:
            raise ValueError("annotation rejected")
        value = await self.fallback.annotate(request)
        return value.model_copy(update={"annotator_id": self.annotator_id})


@pytest.mark.asyncio
async def test_failed_compilation_resumes_from_cached_annotations(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.pptx"
    _make_template(source, slides=2)
    template_dir = tmp_path / "templates" / "flaky"
    annotator = _FlakyAnnotator()
    compiler = TemplateCompiler(renderer=NullRenderer(), annotator=annotator)

    with pytest.raises(TemplateCompilationError, match="annotation rejected"):
        await compiler.compile("flaky", source, template_dir)
    assert (template_dir / ".annotation-cache").exists()

    annotator.fail_slide = None
    manifest = await compiler.compile("flaky", source, template_dir)

    assert manifest.status.value == "ready"
    assert annotator.calls.count("s001") == 1
    assert annotator.calls.count("s002") == 2
    assert not (template_dir / ".annotation-cache").exists()


def test_strict_models_reject_unknown_fields() -> None:
    with pytest.raises(ValidationError):
        NormalizedBox(x=0.0, y=0.0, width=1.0, height=1.0, surprise=True)


def test_revision_id_changes_with_output_affecting_components() -> None:
    common = {
        "source_hash": "a" * 64,
        "compiler_version": "compiler-v1",
        "extractor_version": "extractor-v1",
        "annotator_id": "annotator-v1",
        "renderer_id": "renderer-v1",
    }
    first = derive_revision_id(**common)
    second = derive_revision_id(**{**common, "annotator_id": "annotator-v2"})
    third = derive_revision_id(**{**common, "compiler_version": "compiler-v2"})
    assert first == derive_revision_id(**common)
    assert first != second
    assert first != third


def test_vector_browser_variant_trims_libreoffice_page_canvas(tmp_path: Path) -> None:
    path = tmp_path / "vector-export.png"
    image = Image.new("RGB", (800, 1100), "white")
    image.paste((80, 0, 120), (250, 400, 550, 700))
    image.save(path)

    assert TemplateCompiler._trim_uniform_border(path) is True
    with Image.open(path) as trimmed:
        assert trimmed.width < 320
        assert trimmed.height < 320
