"""Export contracts for shapes PowerPoint has no native equivalent for."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deeppresenter.utils.webview import convert_html_to_pptx

_SLIDE = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0}
.slide{position:relative;width:1280px;height:720px;background:#fff}
.plain{position:absolute;left:5%;top:5%;width:20%;height:20%;background:#005EA4}
.clipped{position:absolute;left:40%;top:5%;width:20%;height:20%;background:#039ACF;
  clip-path:polygon(50% 0%, 100% 100%, 0% 100%)}
</style></head><body><main class="slide">
<div class="plain"></div><div class="clipped"></div>
</main></body></html>"""


def _content_shapes(slide) -> list:
    """Autoshapes the page itself drew, minus the full-canvas background."""

    return [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and not (shape.left == 0 and shape.top == 0 and shape.width >= 12192000)
    ]


def _export(tmp_path: Path) -> Presentation:
    (tmp_path / "slide_01.html").write_text(_SLIDE, encoding="utf-8")
    output = tmp_path / "deck.pptx"
    asyncio.run(convert_html_to_pptx(tmp_path, output, aspect_ratio="16:9"))
    return Presentation(str(output))


def test_a_clipped_shape_survives_the_pptx_export(tmp_path: Path) -> None:
    """PowerPoint cannot express a clip, so the element must be rasterized.

    Exported as an autoshape it becomes the rectangle enclosing the clip, and
    every triangle, chevron and capsule in a template flattens into a block.
    """

    if not Path(__file__).parents[1].joinpath("html2pptx", "node_modules").is_dir():
        pytest.skip("html2pptx dependencies are not installed")

    slide = _export(tmp_path).slides[0]
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    autoshapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]

    # The unclipped block stays a real autoshape; only the clip is rasterized.
    assert len(autoshapes) >= 1
    assert len(pictures) == 1
    picture = pictures[0]
    assert picture.width > 0 and picture.height > 0


_ROTATED = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0}
.slide{position:relative;width:1280px;height:720px;background:#fff}
/* An upward triangle turned upside down, as a template's arrows are. */
.arrow{position:absolute;left:40%;top:30%;width:10%;height:10%;background:#00539E;
  clip-path:polygon(50% 0%, 100% 100%, 0% 100%);transform:rotate(180deg)}
</style></head><body><main class="slide"><div class="arrow"></div></main></body></html>"""


def test_a_rotated_clip_keeps_its_orientation(tmp_path: Path) -> None:
    """getBoundingClientRect reports the rotated box, not the element's own.

    Rasterizing at that size without applying the transform exported every
    arrow unrotated, so a template's downward markers came out as raw
    right-angles pointing the wrong way.
    """

    if not Path(__file__).parents[1].joinpath("html2pptx", "node_modules").is_dir():
        pytest.skip("html2pptx dependencies are not installed")

    import io

    from PIL import Image

    (tmp_path / "slide_01.html").write_text(_ROTATED, encoding="utf-8")
    output = tmp_path / "deck.pptx"
    asyncio.run(convert_html_to_pptx(tmp_path, output, aspect_ratio="16:9"))

    slide = Presentation(str(output)).slides[0]
    picture = next(
        s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    with Image.open(io.BytesIO(picture.image.blob)).convert("RGBA") as image:
        alpha = image.getchannel("A")
        width, height = alpha.size
        top = sum(
            alpha.getpixel((x, y)) > 128
            for y in range(height // 4)
            for x in range(0, width, 4)
        )
        bottom = sum(
            alpha.getpixel((x, y)) > 128
            for y in range(height - height // 4, height)
            for x in range(0, width, 4)
        )

    # A downward triangle is widest at the top.
    assert top > bottom * 2


_STACKED = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0}
.slide{position:relative;width:1280px;height:720px;background:#fff}
/* Written first, painted last: the wedge covers the photo in the browser. */
.wedge{position:absolute;left:10%;top:10%;width:30%;height:30%;
  background:#00539E;z-index:5}
.photo{position:absolute;left:10%;top:10%;width:40%;height:40%;
  background:#CCCCCC;z-index:1}
</style></head><body><main class="slide">
<div class="wedge"></div><div class="photo"></div>
</main></body></html>"""


def test_shapes_are_emitted_in_paint_order_not_document_order(tmp_path: Path) -> None:
    """Document order put a low z-index cover photo on top of the template.

    The photo is written after the wedge but declares a lower z-index, so the
    browser -- and the reference PDF -- paint it underneath. Emitting in
    document order reversed that and buried the template's chrome.
    """

    if not Path(__file__).parents[1].joinpath("html2pptx", "node_modules").is_dir():
        pytest.skip("html2pptx dependencies are not installed")

    (tmp_path / "slide_01.html").write_text(_STACKED, encoding="utf-8")
    output = tmp_path / "deck.pptx"
    asyncio.run(convert_html_to_pptx(tmp_path, output, aspect_ratio="16:9"))

    boxes = _content_shapes(Presentation(str(output)).slides[0])
    # The wider box is the photo; it must come first, i.e. behind the wedge.
    assert len(boxes) == 2
    assert boxes[0].width > boxes[1].width


_OVERHANGING = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0}
.slide{position:relative;width:1280px;height:720px;background:#fff}
/* A brand wedge that bleeds off the left edge, as template chrome does. */
.wedge{position:absolute;left:-400px;top:100px;width:732px;height:400px;
  background:#0071BC;clip-path:polygon(45% 0%, 100% 0%, 55% 100%, 0% 100%)}
</style></head><body><main class="slide"><div class="wedge"></div></main></body></html>"""


def test_decoration_that_overhangs_the_canvas_keeps_its_shape(tmp_path: Path) -> None:
    """Rasterizing at the element's page coordinates clipped to the viewport.

    A wedge starting at a negative x had the part left of the canvas cut, and
    the capture then came back holding the wrong page region -- the artwork
    landed on the wrong side of its own box, so the wedge read as a
    differently shaped block.
    """

    if not Path(__file__).parents[1].joinpath("html2pptx", "node_modules").is_dir():
        pytest.skip("html2pptx dependencies are not installed")

    import io

    from PIL import Image

    (tmp_path / "slide_01.html").write_text(_OVERHANGING, encoding="utf-8")
    output = tmp_path / "deck.pptx"
    asyncio.run(convert_html_to_pptx(tmp_path, output, aspect_ratio="16:9"))

    slide = Presentation(str(output)).slides[0]
    picture = next(
        s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    with Image.open(io.BytesIO(picture.image.blob)).convert("RGBA") as image:
        alpha = image.getchannel("A")
        width, height = alpha.size
        row = height // 2

        def opaque(x: int) -> bool:
            return alpha.getpixel((x, row)) > 128

        # The parallelogram crosses the middle row from about 22% to 78%; the
        # far edges stay empty. Capturing the wrong region filled the left.
        assert not opaque(int(width * 0.05))
        assert opaque(int(width * 0.50))
        assert not opaque(int(width * 0.95))


_BORDERED = """<!doctype html><html><head><meta charset="utf-8"><style>
html,body{margin:0}
.slide{position:relative;width:1280px;height:720px;background:#fff}
.card{position:absolute;left:100px;top:100px;width:400px;height:300px;
  padding:24px;box-sizing:border-box;background:#F5F9FC;
  border-top:8px solid #039ACF}
</style></head><body><main class="slide">
<div class="card"><p>正文</p></div>
</main></body></html>"""


def test_a_border_spans_its_own_box_not_the_padding_box(tmp_path: Path) -> None:
    """A card's accent rule is as wide as the card, padding included.

    Insetting the rule by the element's padding drew every top border short
    of the card it belongs to, so the accent bars floated free of their
    boxes.
    """

    if not Path(__file__).parents[1].joinpath("html2pptx", "node_modules").is_dir():
        pytest.skip("html2pptx dependencies are not installed")

    (tmp_path / "slide_01.html").write_text(_BORDERED, encoding="utf-8")
    output = tmp_path / "deck.pptx"
    asyncio.run(convert_html_to_pptx(tmp_path, output, aspect_ratio="16:9"))

    shapes = _content_shapes(Presentation(str(output)).slides[0])
    card = next(s for s in shapes if s.height > 1000000)
    rule = next(s for s in shapes if s.height <= 1)

    assert rule.left == card.left
    assert rule.width == card.width
